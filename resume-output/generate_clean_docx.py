#!/usr/bin/env python3
"""Clean editable DOCX from original resume content: no dashed boxes, bold keywords."""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = Path(__file__).resolve().parent / "docx-clean"


def set_run_font(run, name_cn="微软雅黑", name_en="Calibri", size=10.5, bold=False, color=None):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = name_en
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), name_en)
    rFonts.set(qn("w:hAnsi"), name_en)
    rFonts.set(qn("w:eastAsia"), name_cn)
    if color is not None:
        run.font.color.rgb = color


def set_para(p, before=0, after=2, line=1.15, align=None):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    if align is not None:
        pf.alignment = align


def add_bottom_border(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "222222")
    pBdr.append(bottom)
    pPr.append(pBdr)


def clear_table_borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else OxmlElement("w:tblPr")
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        el.set(qn("w:sz"), "0")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "auto")
        borders.append(el)
    tblPr.append(borders)


def set_cell_border(cell, **kwargs):
    """kwargs: top/left/bottom/right = (val, sz, color)"""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        if edge in kwargs:
            val, sz, color = kwargs[edge]
            el = OxmlElement(f"w:{edge}")
            el.set(qn("w:val"), val)
            el.set(qn("w:sz"), str(sz))
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), color)
            tcBorders.append(el)
    tcPr.append(tcBorders)


def set_cell_margins(cell, top=40, bottom=40, left=40, right=40):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for m, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        node = OxmlElement(f"w:{m}")
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")
        tcMar.append(node)
    tcPr.append(tcMar)


def add_runs_with_bold(paragraph, parts, size=10, name_cn="微软雅黑", name_en="Calibri"):
    """parts: list of (text, bold)"""
    for text, bold in parts:
        run = paragraph.add_run(text)
        set_run_font(run, name_cn=name_cn, name_en=name_en, size=size, bold=bold)


def bold_keywords(text, keywords, size=10):
    """Split text and bold keyword occurrences. Returns list of (text, bold)."""
    # sort keywords by length desc to match longer first
    keys = sorted(keywords, key=len, reverse=True)
    parts = []
    i = 0
    while i < len(text):
        matched = None
        for k in keys:
            if text.startswith(k, i):
                matched = k
                break
        if matched:
            parts.append((matched, True))
            i += len(matched)
        else:
            # advance to next possible keyword
            next_idx = len(text)
            for k in keys:
                j = text.find(k, i + 1)
                if j != -1:
                    next_idx = min(next_idx, j)
            parts.append((text[i:next_idx], False))
            i = next_idx
    return parts or [(text, False)]


def setup_doc(margins=(1.3, 1.3, 1.5, 1.5)):
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(margins[0])
        section.bottom_margin = Cm(margins[1])
        section.left_margin = Cm(margins[2])
        section.right_margin = Cm(margins[3])
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
    return doc


def add_header_with_photo(doc, lines, photo_label="照片", photo_w=2.2, photo_h=2.8):
    """Header: left text block + right compact photo slot. No outer dashed boxes."""
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    clear_table_borders(table)

    left, right = table.rows[0].cells
    left.width = Cm(14.8)
    right.width = Cm(photo_w + 0.2)

    # clear default empty para and fill left
    left.text = ""
    for i, (text, size, bold) in enumerate(lines):
        p = left.paragraphs[0] if i == 0 else left.add_paragraph()
        set_para(p, before=0 if i else 0, after=2, align=WD_ALIGN_PARAGRAPH.CENTER, line=1.2)
        run = p.add_run(text)
        set_run_font(run, size=size, bold=bold)

    # photo cell
    right.text = ""
    set_cell_margins(right, top=0, bottom=0, left=40, right=0)
    set_cell_border(
        right,
        top=("single", 8, "999999"),
        left=("single", 8, "999999"),
        bottom=("single", 8, "999999"),
        right=("single", 8, "999999"),
    )
    # fixed height via trHeight
    tr = table.rows[0]._tr
    trPr = tr.get_or_add_trPr()
    trHeight = OxmlElement("w:trHeight")
    trHeight.set(qn("w:val"), str(int(photo_h * 567)))  # cm to twips approx
    trHeight.set(qn("w:hRule"), "exact")
    trPr.append(trHeight)

    p = right.paragraphs[0]
    set_para(p, before=28, after=0, align=WD_ALIGN_PARAGRAPH.CENTER)
    run = p.add_run(photo_label)
    set_run_font(run, size=9, color=RGBColor(0x99, 0x99, 0x99))

    # spacer after header
    sp = doc.add_paragraph()
    set_para(sp, before=0, after=2)
    return table


def add_section(doc, title):
    p = doc.add_paragraph()
    set_para(p, before=8, after=3, line=1.1)
    run = p.add_run(title)
    set_run_font(run, size=12, bold=True)
    add_bottom_border(p)


def add_job_head(doc, left, right, size=10.5):
    p = doc.add_paragraph()
    set_para(p, before=5, after=1)
    r1 = p.add_run(left)
    set_run_font(r1, size=size, bold=True)
    r2 = p.add_run("\t" + right)
    set_run_font(r2, size=size - 0.5, bold=False, color=RGBColor(0x55, 0x55, 0x55))
    p.paragraph_format.tab_stops.add_tab_stop(Cm(16.2), WD_TAB_ALIGNMENT.RIGHT)


def add_muted(doc, text, size=9.5):
    p = doc.add_paragraph()
    set_para(p, before=0, after=1)
    run = p.add_run(text)
    set_run_font(run, size=size, color=RGBColor(0x55, 0x55, 0x55))


def add_bullet_rich(doc, text, keywords, size=9.8, mark="❖ "):
    p = doc.add_paragraph()
    set_para(p, before=0, after=1.5, line=1.2)
    run0 = p.add_run(mark)
    set_run_font(run0, size=size)
    for frag, bold in bold_keywords(text, keywords):
        run = p.add_run(frag)
        set_run_font(run, size=size, bold=bold)


def add_line_rich(doc, text, keywords, size=9.8):
    p = doc.add_paragraph()
    set_para(p, before=0, after=1.5, line=1.2)
    for frag, bold in bold_keywords(text, keywords):
        run = p.add_run(frag)
        set_run_font(run, size=size, bold=bold)


ZH_KEYS = [
    "海外技术支持工程师",
    "海外技术支持",
    "服务工程师",
    "半导体设备",
    "机器人",
    "复合岗",
    "复合岗位",
    "西班牙",
    "荷兰",
    "法国",
    "德国",
    "四轮定位",
    "ADAS",
    "Cursor",
    "AI Agent",
    "AI",
    "SmartSafe",
    "易检智能",
    "ELEGOO",
    "PCB",
    "STM32",
    "FreeRTOS",
    "UART",
    "SPI",
    "GPS",
    "IMU",
    "PID",
    "μC/OS-II",
    "CET-6",
    "英国签证",
    "深圳",
    "知识库",
    "工单",
    "售前培训",
    "售后",
    "研发",
    "智能车实验室",
    "英飞凌",
    "TC264",
    "TC377",
    "MSPM0G3507",
    "MPU6050",
    "Python",
    "Matlab",
    "浙江省二等奖",
    "校一等奖",
    "全国三等奖",
]

EN_KEYS = [
    "Overseas Technical Support",
    "Service Engineer",
    "Semiconductor Equipment",
    "Robotics",
    "Hybrid",
    "SmartSafe",
    "ELEGOO",
    "Spain",
    "Netherlands",
    "France",
    "Germany",
    "wheel-alignment",
    "wheel alignment",
    "ADAS",
    "Cursor",
    "AI",
    "R&D",
    "UK visa",
    "Shenzhen",
    "STM32",
    "FreeRTOS",
    "Infineon",
    "TC264",
    "TC377",
    "GPS",
    "IMU",
    "PID",
    "μC/OS-II",
    "CET-6",
    "knowledge base",
    "knowledge-base",
    "ticket",
    "Provincial 2nd Prize",
]


def build_zh_1():
    doc = setup_doc((1.2, 1.2, 1.5, 1.5))
    add_header_with_photo(
        doc,
        [
            ("侯钟棋", 20, True),
            ("求职意向：海外技术支持工程师（半导体设备 / 机器人 · 复合岗）｜深圳｜可接受海外出差", 10.5, True),
            ("23岁｜男｜本科｜英语可流利对接海外客户｜持有英国签证（约2年有效）", 10, False),
            ("178-1676-9251｜zhongqi_hou@foxmail.com", 10, False),
        ],
    )

    add_section(doc, "核心优势")
    for t in [
        "海外客户现场与远程支持：欧洲多国（西班牙 / 荷兰 / 法国 / 德国）售后处理与产品培训，英文听说沟通基本无口音。",
        "客户—研发桥梁：跟进大客户四轮定位设备软硬件需求约4个月，推动需求澄清、方案对齐与产品落地。",
        "工具化提效：用 Cursor 搭建服务知识库智能助手（公司服务器 7×24 运行），并开发售后工单与数据看板平台。",
        "技术底座：自动化专业 + 智能车实验室，具备 PCB 识图、嵌入式 C、焊接与万用表等基础，学习新产品线速度快。",
    ]:
        add_bullet_rich(doc, t, ZH_KEYS)

    add_section(doc, "工作经历")
    add_job_head(doc, "易检智能（SmartSafe）｜海外技术支持 / 服务工程师", "2025.11 – 至今")
    add_muted(doc, "汽车后市场检测与维保设备｜负责售前培训、售后支持及客户与研发之间的技术沟通")
    for t in [
        "覆盖产品线：四轮定位、ADAS 校准、冷媒交换机、喷油嘴检测清洗、变速箱油液交换、新能源电池包充放电/均衡维保、诊断平板等；快速完成多线产品上手并独立支持客户。",
        "海外出差：赴西班牙、荷兰、法国、德国开展售后问题处理与现场产品培训，面向经销商/终端客户完成演示、答疑与故障闭环。",
        "大客户交付：主导四轮定位设备软硬件需求对接约4个月，在客户与研发之间拆解需求、对齐方案与验证节奏，最终推动产品落地。",
        "智能知识库：用 Cursor 搭建服务部门问答式知识库（问题处理办法 + 产品文档），为销售/服务同事开通账号并部署于公司服务器 7×24 运行，缓解时差导致的资料检索瓶颈。",
        "工单平台：开发售后工单记录系统，支持视频/图片/文档上传与客户问题看板，实现 case 可追溯与问题累计可视化。",
    ]:
        add_bullet_rich(doc, t, ZH_KEYS)

    add_job_head(doc, "ELEGOO（深圳市智能派科技有限公司）｜海外技术支持工程师", "2025.08 – 2025.10")
    add_muted(doc, "3D 打印出海｜远程服务欧美客户技术咨询与问题闭环")
    for t in [
        "快速掌握产品线，独立处理欧美客户技术问题；整理典型案例与解决方案，搭建个人知识库并训练 AI Agent 提升响应效率。",
        "将一线高频问题、设计缺陷与改进建议系统反馈产品/技术部门，形成可复用的迭代输入。",
    ]:
        add_bullet_rich(doc, t, ZH_KEYS)

    add_section(doc, "教育背景")
    add_job_head(doc, "浙江科技大学｜自动化｜本科｜GPA 3.04（专业前15%）", "2021.09 – 2025.06")
    for t in [
        "智能车实验室 3 年：英飞凌 TC264/TC377 惯导定位越野车（GPS+IMU 融合、PID 整定），获第十八届全国大学生智能车竞赛浙江省二等奖（省第四）。",
        "具备 STM32 / 传感器驱动 / UART·SPI / FreeRTOS 基础，以及 PCB 识图、焊接、万用表调试与嵌入式 C 实践经验。",
    ]:
        add_bullet_rich(doc, t, ZH_KEYS)

    add_section(doc, "技能特长")
    for t in [
        "语言：英语 CET-6；可独立完成海外客户技术沟通、现场培训与英文邮件/文档处理。",
        "支持能力：多产品线售后诊断、售前培训、需求转译、跨部门协调（客户 ↔ 研发 ↔ 销售）。",
        "技术基础：嵌入式与自动化背景；Python/Matlab 数据分析；Cursor / AI 工具提效；流程文档与知识库建设。",
        "出差意愿：约 1 年 3 次，单次约 1 个月；工作地期望深圳。",
    ]:
        add_line_rich(doc, t, ZH_KEYS)

    add_section(doc, "奖项荣誉")
    for t in [
        "第十八届全国大学生智能车竞赛 浙江省二等奖（2023.07）｜浙江省大学生英语竞赛 校一等奖（2022.05）",
        "第六届普译奖全国大学生英语写作大赛 全国三等奖（2023.12）｜学业优秀奖学金 / 美仪企业奖学金（7/92）",
    ]:
        add_bullet_rich(doc, t, ZH_KEYS)

    add_section(doc, "自我评价")
    add_line_rich(
        doc,
        "沟通协调强、学习速度快，擅长在客户现场与研发之间把问题说清楚、把方案推落地。目标先深耕海外技术支持，积累设备与行业经验后向海外销售/复合型客户成功路径发展。",
        ZH_KEYS,
    )
    return doc


def build_zh_2():
    doc = setup_doc((1.4, 1.4, 1.5, 1.5))
    add_header_with_photo(
        doc,
        [
            ("侯钟棋", 20, True),
            ("求职意向：海外技术支持工程师（半导体设备 / 机器人 · 复合岗）｜工作地：深圳", 10.5, True),
            ("23岁｜男｜自动化本科｜英语可流利对接海外客户（基本无口音）｜持有英国签证（约2年有效）", 10, False),
            ("电话：178-1676-9251｜邮箱：zhongqi_hou@foxmail.com｜出差意愿：约1年3次，单次约1个月", 10, False),
        ],
    )

    add_section(doc, "个人简介")
    for t in [
        "目标路径：先入行海外技术支持，深耕设备理解与客户信任，再向海外销售 / 售前复合岗位发展。",
        "当前价值：能独立服务海外客户（远程 + 欧洲现场），擅长把客户现场问题翻译成研发可执行需求，并把方案再翻译回客户可落地动作。",
        "效率杠杆：善用 AI（Cursor）建设部门级知识库与售后工单系统，解决跨时区资料检索与 case 沉淀问题。",
        "技术底座：自动化专业 + 智能车实验室，具备传感器、嵌入式、PCB 与现场调试基础，适合机电/检测/机器人相关设备支持岗位。",
    ]:
        add_bullet_rich(doc, t, ZH_KEYS)

    add_section(doc, "工作经历")
    add_job_head(doc, "易检智能（SmartSafe）｜海外技术支持 / 服务工程师", "2025.11 – 至今")
    add_muted(doc, "行业：汽车后市场检测与维保设备｜角色：对内服务工程师，对外海外技术支持；售前培训 + 售后支持 + 客户/研发桥梁")
    for t in [
        "多产品线支持：四轮定位设备、ADAS 校准设备、汽车空调冷媒交换机、喷油嘴检测清洗机、变速箱油液清洗交换机、新能源电池包充放电与均衡维保设备、汽车诊断平板等；在培训深度有限的情况下，依靠自学与实战快速形成可交付支持能力。",
        "海外现场：出差西班牙、荷兰、法国、德国，完成售后故障处理、产品培训与客户答疑，积累跨文化沟通与现场闭环经验。",
        "大客户项目：跟随大客户四轮定位设备软硬件需求约4个月，持续在研发与客户两侧对齐功能、接口、验收与问题优先级，最终推动产品落地交付。",
        "角色定位：技术深度因产品线广而不追求单机专家，但在需求澄清、问题定位、跨部门推进与客户预期管理方面形成稳定输出。",
        "AI 知识库助手：使用 Cursor 搭建服务部门智能知识库，沉淀常见问题、处理办法与产品文件；同事以对话方式提问即可定位对应文档；为销售与服务同事开通账号，部署于公司服务器并 7×24 运行，缓解国内外时差导致“销售要资料、国内已下班”的效率损失。",
        "售后工单平台：开发工单记录系统，支持上传视频、图片、文档，按 case 关联客户与问题；配套数据看板统计客户问题累计，提升售后过程透明度与复盘效率。",
    ]:
        add_bullet_rich(doc, t, ZH_KEYS)

    add_job_head(doc, "ELEGOO（深圳市智能派科技有限公司）｜海外技术支持工程师", "2025.08 – 2025.10")
    add_muted(doc, "行业：3D 打印出海｜工作模式：工位远程支持欧美客户售后邮件/技术咨询")
    for t in [
        "快速学习公司产品线，独立处理欧美客户技术咨询，日均处理大量问题并保持高好评反馈；形成“先复现、再分层、再闭环”的支持习惯。",
        "沉淀个人知识库并训练 AI Agent，将重复问题标准化，提高检索与回复效率。",
        "系统整理用户痛点、设计缺陷与改进建议，反馈产品与技术部门，作为迭代输入。",
    ]:
        add_bullet_rich(doc, t, ZH_KEYS)

    add_section(doc, "教育背景")
    add_job_head(doc, "浙江科技大学｜自动化｜本科｜GPA 3.04（专业前15%）", "2021.09 – 2025.06")
    add_muted(doc, "英语水平：CET-6｜主修：单片机与嵌入式、微机原理、自动控制原理、C 语言、数据结构等｜智能车实验室 3 年")

    add_section(doc, "项目经历（技术基础）")
    add_job_head(doc, "惯导定位越野车（英飞凌 TC264 / TC377）", "2022.12 – 2023.08")
    add_muted(doc, "智能车实验室｜感知：GPS + IMU｜成果：浙江省大学生智能车竞赛二等奖（省第四）")
    for t in [
        "对比 GPS_TAU1201（UART）与 IMU660RA（SPI）方案，采用融合与互补算法提升实时性，缩短循迹响应，减少约 2s 比赛用时。",
        "设计自动任务分区逻辑，降低人工干预；用 VOFA+ 与 Matlab 分析路径误差并优化 PID，降低寻迹偏移。",
        "备赛中协调传感器选型与软硬件分工，锻炼跨角色协作与问题决策能力。",
    ]:
        add_bullet_rich(doc, t, ZH_KEYS)

    add_job_head(doc, "自动红外循迹系统（德州仪器 MSPM0G3507）", "2024.07 – 2024.08")
    add_muted(doc, "MPU6050 + 红外阵列｜成果：浙江省大学生电子设计竞赛 H 题（完成前三问）")
    for t in [
        "设计 7 路红外循迹阵列（3-1-3）布局并推动团队落地；用 DMP 获取欧拉角并建立误差修正，缓解偏航角零漂。",
        "基于 Keil5 / SysConfig 配置 GPIO/UART，缩短外设联调周期。",
    ]:
        add_bullet_rich(doc, t, ZH_KEYS)

    add_job_head(doc, "基于 μC/OS-II 车载电控单元（STM32F103C8T6）", "2025.02 – 2025.04")
    add_muted(doc, "实时控制与安全响应｜成果：障碍物识别响应延迟 <50ms")
    for t in [
        "对比 FreeRTOS 与 μC/OS-II 后完成系统移植；主导超声波测距、电机减速、蜂鸣报警与霍尔转速可视化逻辑。",
        "使用嘉立创 EDA 迭代 PCB，缩短硬件验证周期；形成“指标定义—方案选型—联调验证”的完整闭环习惯。",
    ]:
        add_bullet_rich(doc, t, ZH_KEYS)

    add_section(doc, "技能特长")
    for t in [
        "语言能力：英语 CET-6；可独立完成海外客户口语沟通、现场培训讲解与英文技术邮件处理。",
        "客户支持：售前培训、售后诊断、现场交付、需求转译、跨时区沟通、经销商赋能。",
        "工程基础：STM32 / 英飞凌 / TI 单片机；UART/SPI；FreeRTOS / μC/OS-II 基础；PCB 识图、焊接、万用表；Python（Matplotlib / Numpy / Pandas）、Matlab。",
        "效率工具：Cursor / AI Agent；知识库建设；工单与数据看板；ProcessOn / Visio / Xmind / Office。",
        "其他：持有英国签证（约2年有效）；期望工作地深圳；可接受海外出差（约1年3次，单次约1个月）。",
    ]:
        add_line_rich(doc, t, ZH_KEYS)

    add_section(doc, "奖项荣誉")
    for t in [
        "第十八届全国大学生智能车竞赛 浙江省二等奖（2023.07）",
        "2024 年浙江省大学生电子设计竞赛 参赛完成（H 题前三问，2024.08）",
        "浙江省大学生英语竞赛 校一等奖（2022.05）",
        "第六届普译奖全国大学生英语写作大赛 全国三等奖（2023.12）",
        "学业优秀一等奖学金 / 美仪企业奖学金风华学子电气拓新奖（7/92，2022–2024）",
    ]:
        add_bullet_rich(doc, t, ZH_KEYS)

    add_section(doc, "自我评价")
    for t in [
        "沟通能力是当前最强项：能在客户、销售、研发之间快速对齐事实与下一步动作，推动问题闭环。",
        "学习与工具化能力强：面对多产品线与有限培训，能靠自学、文档沉淀和 AI 提效形成可复制支持能力。",
        "编程能力仍在持续提升；更适合以“设备理解 + 客户交付 + 跨部门推进”为核心的海外技术支持 / 售前复合岗位。",
        "性格积极，院篮球队经历；希望加入半导体设备或机器人行业，在深圳长期发展，并逐步走向海外销售。",
    ]:
        add_bullet_rich(doc, t, ZH_KEYS)
    return doc


def build_en_1():
    doc = setup_doc((1.2, 1.2, 1.5, 1.5))
    add_header_with_photo(
        doc,
        [
            ("Zhongqi Hou", 20, True),
            ("Overseas Technical Support Engineer (Semiconductor Equipment / Robotics · Hybrid) | Shenzhen", 10, True),
            ("English-fluent | UK visa (~2 years) | Overseas travel OK (~3×/year, ~1 month) | +86 178-1676-9251 | zhongqi_hou@foxmail.com", 9.5, False),
        ],
        photo_label="Photo",
    )

    add_section(doc, "SUMMARY")
    for t in [
        "Europe field support (Spain / Netherlands / France / Germany) + remote overseas customers; clear spoken English for training and troubleshooting.",
        "Customer–R&D bridge: drove key-account wheel-alignment HW/SW requirements for ~4 months through to product delivery.",
        "Built AI knowledge-base assistant (Cursor, 24/7 on company servers) and ticket/dashboard platform for after-sales efficiency across time zones.",
        "Automation degree + smart-car lab (PCB, embedded C, sensors)—fast ramp-up on electromechanical equipment.",
    ]:
        add_bullet_rich(doc, t, EN_KEYS, mark="• ")

    add_section(doc, "EXPERIENCE")
    add_job_head(doc, "SmartSafe (Yijian Intelligent) — Overseas Technical Support / Service Engineer", "Nov 2025 – Present")
    add_muted(doc, "Automotive aftermarket test & maintenance equipment | Pre-sales training, after-sales, customer–R&D coordination")
    for t in [
        "Multi-line support: wheel alignment, ADAS calibration, A/C refrigerant exchange, injector cleaning, ATF exchange, EV battery pack charge/discharge & balancing, diagnostic tablets.",
        "On-site in ES/NL/FR/DE for fault handling, product training, and dealer/end-user enablement.",
        "Led ~4-month key-account wheel-alignment HW/SW alignment between customer and R&D until delivery landed.",
        "Built department AI knowledge base (FAQ + docs, chat Q&A) with Cursor; 24/7 server deploy for sales/service; built ticket platform with media upload + issue dashboard.",
    ]:
        add_bullet_rich(doc, t, EN_KEYS, mark="• ")

    add_job_head(doc, "ELEGOO (Shenzhen Zhinenpai Technology) — Overseas Technical Support Engineer", "Aug 2025 – Oct 2025")
    add_muted(doc, "3D printing export | Remote EU/US technical support")
    add_bullet_rich(
        doc,
        "Independent EU/US technical inquiry handling; personal knowledge base + AI agent for faster response; structured feedback to product/engineering.",
        EN_KEYS,
        mark="• ",
    )

    add_section(doc, "EDUCATION")
    add_job_head(doc, "Zhejiang University of Science and Technology — B.E. Automation | GPA 3.04 (Top 15%)", "Sep 2021 – Jun 2025")
    for t in [
        "Intelligent Vehicle Lab (3 yrs): Infineon TC264/TC377 GPS+IMU vehicle; Provincial 2nd Prize, National Intelligent Vehicle Competition (4th in Zhejiang).",
        "STM32 / sensor drivers / UART·SPI / FreeRTOS basics; PCB reading, soldering, multimeter, embedded C.",
    ]:
        add_bullet_rich(doc, t, EN_KEYS, mark="• ")

    add_section(doc, "SKILLS & AWARDS")
    for t in [
        "Languages: CET-6 English — overseas calls, on-site training, technical email.",
        "Support: multi-product diagnosis, pre-sales training, requirement translation, customer ↔ R&D ↔ sales coordination.",
        "Technical: automation/embedded base; Python/Matlab; Cursor/AI tooling; knowledge-base & ticket workflows. Prefer Shenzhen base.",
        "Awards: Provincial 2nd Prize, National Intelligent Vehicle Competition (2023); Campus 1st, Zhejiang College English Contest (2022); National 3rd, Puyi English Writing (2023); scholarships (7/92).",
    ]:
        add_line_rich(doc, t, EN_KEYS)

    add_section(doc, "PROFILE")
    add_line_rich(
        doc,
        "Strong communicator and fast learner; turns field issues into actionable R&D inputs and customer-ready solutions. Aiming to deepen overseas technical support in semiconductor equipment or robotics, then grow toward overseas sales / hybrid roles.",
        EN_KEYS,
    )
    return doc


def build_en_2():
    doc = setup_doc((1.4, 1.4, 1.5, 1.5))
    add_header_with_photo(
        doc,
        [
            ("Zhongqi Hou", 20, True),
            ("Target: Overseas Technical Support Engineer (Semiconductor Equipment / Robotics · Hybrid)", 10.5, True),
            ("Shenzhen, China | English-fluent with overseas customers | UK visa holder (~2 years)", 10, False),
            ("+86 178-1676-9251 | zhongqi_hou@foxmail.com | Travel: ~3 overseas trips/year, ~1 month each", 10, False),
        ],
        photo_label="Photo",
    )

    add_section(doc, "PROFESSIONAL SUMMARY")
    for t in [
        "Path: deepen overseas technical support first, then move toward overseas sales / pre-sales hybrid roles.",
        "Strength: independent overseas delivery (remote + Europe on-site); translate field problems into R&D-executable requirements and customer-actionable steps.",
        "Efficiency: use AI (Cursor) to build department knowledge systems and ticket tooling that cut timezone friction.",
        "Base: Automation degree + Intelligent Vehicle Lab; sensor/embedded/PCB/debug foundation for electromechanical and robotics-related equipment support.",
    ]:
        add_bullet_rich(doc, t, EN_KEYS, mark="• ")

    add_section(doc, "WORK EXPERIENCE")
    add_job_head(doc, "SmartSafe (Yijian Intelligent) — Overseas Technical Support / Service Engineer", "Nov 2025 – Present")
    add_muted(doc, "Automotive aftermarket inspection & maintenance equipment | Internal service engineer + external overseas technical support")
    for t in [
        "Multi-line support: wheel alignment, ADAS calibration, A/C refrigerant exchange, injector cleaning, ATF exchange, EV battery pack charge/discharge & balancing, diagnostic tablets—ramped up quickly under limited deep training.",
        "Field work in Spain, Netherlands, France, and Germany: after-sales troubleshooting, product training, dealer/end-user enablement.",
        "Key account: followed wheel-alignment HW/SW requirements for ~4 months; aligned features, interfaces, acceptance, and priorities between R&D and customer until product landed.",
        "Role focus: reliable at requirement clarification, issue triage, cross-team push, and expectation management across a wide portfolio.",
        "AI knowledge base (Cursor): FAQ + product docs, chat-style retrieval; accounts for sales/service; 24/7 on company servers to reduce after-hours document search across time zones.",
        "Ticket platform: case records with video/image/document upload, customer–issue linkage, and dashboard for cumulative issue visibility.",
    ]:
        add_bullet_rich(doc, t, EN_KEYS, mark="• ")

    add_job_head(doc, "ELEGOO (Shenzhen Zhinenpai Technology) — Overseas Technical Support Engineer", "Aug 2025 – Oct 2025")
    add_muted(doc, "3D printing export | Desk-based remote after-sales/technical support for EU/US customers")
    for t in [
        "Independent EU/US technical inquiries at high daily volume; built personal knowledge base + AI agent to standardize recurring issues.",
        "Compiled user pain points and design gaps as structured input for product and engineering teams.",
    ]:
        add_bullet_rich(doc, t, EN_KEYS, mark="• ")

    add_section(doc, "EDUCATION")
    add_job_head(doc, "Zhejiang University of Science and Technology — B.E. Automation | GPA 3.04 (Top 15%)", "Sep 2021 – Jun 2025")
    add_muted(doc, "CET-6 | Coursework: MCU/embedded, microcomputer principles, automatic control, C, data structures | Intelligent Vehicle Lab 3 years")

    add_section(doc, "PROJECTS (TECHNICAL FOUNDATION)")
    add_job_head(doc, "Inertial Navigation Off-Road Vehicle (Infineon TC264 / TC377)", "Dec 2022 – Aug 2023")
    add_muted(doc, "GPS + IMU | Provincial 2nd Prize (4th in Zhejiang), National Intelligent Vehicle Competition")
    for t in [
        "Compared GPS (UART) vs IMU (SPI); fusion/complementary filtering improved real-time response and cut ~2s race time.",
        "Automatic task-zone logic; VOFA+ / Matlab path-error analysis and PID tuning; coordinated sensor selection and HW/SW split.",
    ]:
        add_bullet_rich(doc, t, EN_KEYS, mark="• ")

    add_job_head(doc, "Infrared Line-Following (TI MSPM0G3507) · Vehicle ECU on μC/OS-II (STM32F103)", "2024.07 – 2025.04")
    for t in [
        "IR array + MPU6050 with DMP Euler angles and yaw-drift correction; Zhejiang Electronic Design Contest Problem H (first three tasks).",
        "Ported μC/OS-II; ultrasonic ranging, motor deceleration, buzzer, Hall RPM display; obstacle response <50ms; PCB iteration in JLCPCB EDA.",
    ]:
        add_bullet_rich(doc, t, EN_KEYS, mark="• ")

    add_section(doc, "SKILLS")
    for t in [
        "Languages: CET-6 English; overseas customer conversations, on-site training, technical email.",
        "Support: pre-sales training, after-sales diagnosis, field delivery, requirement translation, cross-timezone communication.",
        "Engineering: STM32 / Infineon / TI; UART/SPI; RTOS basics; PCB reading, soldering, multimeter; Python / Matlab.",
        "Productivity: Cursor / AI agents; knowledge-base building; ticket & dashboard workflows; ProcessOn / Visio / Xmind / Office.",
        "Other: UK visa (~2 years); Shenzhen base preferred; overseas travel OK (~3×/year, ~1 month each).",
    ]:
        add_line_rich(doc, t, EN_KEYS)

    add_section(doc, "HONORS")
    for t in [
        "Provincial 2nd Prize, 18th National College Student Intelligent Vehicle Competition (2023.07)",
        "Campus 1st Prize, Zhejiang College English Contest (2022.05) · National 3rd Prize, Puyi English Writing Contest (2023.12)",
        "Academic Excellence / Meiyi Corporate Scholarship (rank 7/92, 2022–2024)",
    ]:
        add_bullet_rich(doc, t, EN_KEYS, mark="• ")

    add_section(doc, "SELF ASSESSMENT")
    add_line_rich(
        doc,
        "Communication is the strongest asset—align facts and next actions among customers, sales, and R&D until issues close. Fast learner with a tooling mindset. Coding depth is still growing; best fit is overseas technical support / pre-sales hybrid roles in semiconductor equipment or robotics, with a medium-term path toward overseas sales in Shenzhen.",
        EN_KEYS,
    )
    return doc


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    files = {
        "侯钟棋-海外技术支持-中文-一页版.docx": build_zh_1(),
        "侯钟棋-海外技术支持-中文-两页版.docx": build_zh_2(),
        "Hou-Zhongqi-Overseas-Technical-Support-EN-1page.docx": build_en_1(),
        "Hou-Zhongqi-Overseas-Technical-Support-EN-2page.docx": build_en_2(),
    }
    for name, doc in files.items():
        path = OUT / name
        doc.save(path)
        print("Wrote", path)


if __name__ == "__main__":
    main()
