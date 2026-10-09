#!/usr/bin/env python3
"""Generate four professional .docx resumes (no hybrid labels, no semicolons)."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = Path(__file__).resolve().parent / "docx"


def set_run_font(run, name_cn="微软雅黑", name_en="Calibri", size=10.5, bold=False, color=None):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = name_en
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name_cn)
    if color is not None:
        run.font.color.rgb = color


def set_para(p, before=0, after=2, line=1.15, align=None):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    if align is not None:
        pf.alignment = align


def add_heading_bar(doc, text):
    p = doc.add_paragraph()
    set_para(p, before=8, after=2, line=1.1)
    run = p.add_run(text)
    set_run_font(run, size=12, bold=True)
    pPr = p._p.get_or_add_pPr()
    pBdr = pPr.makeelement(qn("w:pBdr"), {})
    bottom = pBdr.makeelement(
        qn("w:bottom"),
        {
            qn("w:val"): "single",
            qn("w:sz"): "12",
            qn("w:space"): "1",
            qn("w:color"): "222222",
        },
    )
    pBdr.append(bottom)
    pPr.append(pBdr)


def add_center(doc, text, size=10.5, bold=False, before=0, after=2):
    p = doc.add_paragraph()
    set_para(p, before=before, after=after, align=WD_ALIGN_PARAGRAPH.CENTER)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)
    return p


def add_line(doc, text, size=10.5, bold=False, before=0, after=2, color=None):
    p = doc.add_paragraph()
    set_para(p, before=before, after=after)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    return p


def add_job_head(doc, left, right, size=10.5):
    p = doc.add_paragraph()
    set_para(p, before=5, after=1)
    r1 = p.add_run(left)
    set_run_font(r1, size=size, bold=True)
    r2 = p.add_run("\t" + right)
    set_run_font(r2, size=size - 0.5, bold=False, color=RGBColor(0x55, 0x55, 0x55))
    p.paragraph_format.tab_stops.add_tab_stop(Cm(16.5), WD_TAB_ALIGNMENT.RIGHT)
    return p


def add_bullet(doc, text, size=10):
    p = doc.add_paragraph(style="List Bullet")
    set_para(p, before=0, after=1, line=1.2)
    p.clear()
    run = p.add_run(text)
    set_run_font(run, size=size)
    return p


def add_muted(doc, text, size=9.5):
    return add_line(doc, text, size=size, before=0, after=1, color=RGBColor(0x55, 0x55, 0x55))


def setup_doc(margins=(1.4, 1.5, 1.4, 1.5)):
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(margins[0])
        section.bottom_margin = Cm(margins[1])
        section.left_margin = Cm(margins[2])
        section.right_margin = Cm(margins[3])
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
    return doc


def build_zh_1():
    doc = setup_doc((1.2, 1.2, 1.4, 1.4))
    add_center(doc, "侯钟棋", size=20, bold=True, after=4)
    add_center(doc, "求职意向：海外技术支持工程师（半导体设备 / 机器人）｜深圳｜可接受海外出差", size=10.5, bold=True)
    add_center(doc, "23岁｜男｜本科｜英语可流利对接海外客户｜持有英国签证（约2年有效）", size=10)
    add_center(doc, "178-1676-9251｜zhongqi_hou@foxmail.com", size=10, after=4)

    add_heading_bar(doc, "核心优势")
    for t in [
        "具备海外客户现场与远程支持经验，曾赴西班牙、荷兰、法国、德国开展售后处理与产品培训，英语听说沟通顺畅。",
        "擅长在客户与研发之间进行需求对接与进度协同。曾跟进大客户四轮定位设备软硬件需求约4个月，完成需求澄清、方案对齐并推动产品交付。",
        "使用 Cursor 搭建服务部门知识库助手，部署于公司服务器并全天运行。同时开发售后工单记录与问题数据看板，提升跨时区资料检索与 case 管理效率。",
        "自动化专业背景，具备智能车实验室实践经验，掌握 PCB 识图、嵌入式 C、焊接与万用表调试等基础，能够较快完成新产品线上手。",
    ]:
        add_bullet(doc, t, size=9.5)

    add_heading_bar(doc, "工作经历")
    add_job_head(doc, "易检智能（SmartSafe）｜海外技术支持 / 服务工程师", "2025.11 – 至今")
    add_muted(doc, "汽车后市场检测与维保设备｜负责售前培训、售后支持及客户与研发之间的技术沟通")
    for t in [
        "负责产品线包括四轮定位、ADAS 校准、冷媒交换机、喷油嘴检测清洗、变速箱油液交换、新能源电池包充放电与均衡维保、诊断平板等。在培训深度有限的情况下，通过自学与实战快速形成可独立支持客户的能力。",
        "赴西班牙、荷兰、法国、德国开展售后问题处理与现场产品培训，面向经销商与终端客户完成演示、答疑与故障处理。",
        "主导大客户四轮定位设备软硬件需求对接约4个月，在客户与研发之间拆解需求、对齐方案与验证节奏，最终完成产品交付。",
        "使用 Cursor 搭建服务部门问答式知识库，沉淀常见问题、处理办法与产品文档，为销售与服务同事开通账号，并部署于公司服务器全天运行，缓解时差导致的资料检索不便。",
        "开发售后工单记录系统，支持视频、图片与文档上传，按 case 关联客户与问题，并配套数据看板统计客户问题累计情况。",
    ]:
        add_bullet(doc, t, size=9.5)

    add_job_head(doc, "ELEGOO（深圳市智能派科技有限公司）｜海外技术支持工程师", "2025.08 – 2025.10")
    add_muted(doc, "3D 打印出海｜远程处理欧美客户技术咨询与问题跟进")
    for t in [
        "快速掌握公司产品线，独立处理欧美客户技术问题。整理典型案例与解决方案，搭建个人知识库并训练 AI Agent，提升响应效率。",
        "将一线高频问题、设计缺陷与改进建议整理反馈至产品与技术部门，作为产品迭代输入。",
    ]:
        add_bullet(doc, t, size=9.5)

    add_heading_bar(doc, "教育背景")
    add_job_head(doc, "浙江科技大学｜自动化｜本科｜GPA 3.04（专业前15%）", "2021.09 – 2025.06")
    for t in [
        "智能车实验室三年。参与英飞凌 TC264/TC377 惯导定位越野车项目，完成 GPS 与 IMU 融合及 PID 整定，获第十八届全国大学生智能车竞赛浙江省二等奖（省第四）。",
        "具备 STM32、传感器驱动、UART/SPI、FreeRTOS 基础，以及 PCB 识图、焊接、万用表调试与嵌入式 C 实践经验。",
    ]:
        add_bullet(doc, t, size=9.5)

    add_heading_bar(doc, "技能特长")
    for t in [
        "语言：英语 CET-6，可独立完成海外客户技术沟通、现场培训与英文邮件、文档处理。",
        "支持能力：多产品线售后诊断、售前培训、需求转译、跨部门协调（客户、研发、销售）。",
        "技术基础：嵌入式与自动化背景，具备 Python/Matlab 数据分析能力，熟悉 Cursor 等 AI 工具提效，以及流程文档与知识库建设。",
        "其他：期望工作地深圳。可接受海外出差，约一年三次，单次约一个月。",
    ]:
        add_bullet(doc, t, size=9.5)

    add_heading_bar(doc, "奖项荣誉")
    for t in [
        "第十八届全国大学生智能车竞赛 浙江省二等奖（2023.07）",
        "浙江省大学生英语竞赛 校一等奖（2022.05）",
        "第六届普译奖全国大学生英语写作大赛 全国三等奖（2023.12）",
        "学业优秀奖学金、美仪企业奖学金（7/92）",
    ]:
        add_bullet(doc, t, size=9.5)

    add_heading_bar(doc, "自我评价")
    add_line(
        doc,
        "具备较强的沟通协调与快速学习能力，擅长在客户现场与研发之间澄清问题并推进方案落地。希望深耕海外技术支持，积累设备与行业经验后，逐步向海外销售方向发展。",
        size=9.5,
        after=2,
    )
    return doc


def build_zh_2():
    doc = setup_doc((1.4, 1.4, 1.5, 1.5))
    add_center(doc, "侯钟棋", size=20, bold=True, after=4)
    add_center(doc, "求职意向：海外技术支持工程师（半导体设备 / 机器人）｜深圳", size=10.5, bold=True)
    add_center(doc, "23岁｜男｜自动化本科｜英语可流利对接海外客户｜持有英国签证（约2年有效）", size=10)
    add_center(doc, "电话：178-1676-9251｜邮箱：zhongqi_hou@foxmail.com｜出差意愿：约一年三次，单次约一个月", size=10, after=4)

    add_heading_bar(doc, "个人简介")
    for t in [
        "职业规划：先深耕海外技术支持，夯实设备理解与客户信任，再向海外销售或售前方向发展。",
        "当前优势：可独立承担海外客户远程与欧洲现场支持，能够将客户现场问题转化为研发可执行需求，并将方案以客户可理解的方式同步落地。",
        "效率能力：善用 Cursor 建设部门级知识库与售后工单系统，提升跨时区资料检索与 case 沉淀效率。",
        "技术基础：自动化专业，具备智能车实验室传感器、嵌入式、PCB 与现场调试经验，适合机电、检测及机器人相关设备支持岗位。",
    ]:
        add_bullet(doc, t)

    add_heading_bar(doc, "工作经历")
    add_job_head(doc, "易检智能（SmartSafe）｜海外技术支持 / 服务工程师", "2025.11 – 至今")
    add_muted(doc, "汽车后市场检测与维保设备｜对内服务工程师，对外海外技术支持，覆盖售前培训与售后支持")
    for t in [
        "负责产品线包括四轮定位、ADAS 校准、汽车空调冷媒交换机、喷油嘴检测清洗机、变速箱油液清洗交换机、新能源电池包充放电与均衡维保设备、汽车诊断平板等。在培训深度有限的情况下，通过自学与实战快速形成可交付支持能力。",
        "出差西班牙、荷兰、法国、德国，完成售后故障处理、产品培训与客户答疑，积累跨文化沟通与现场闭环经验。",
        "跟进大客户四轮定位设备软硬件需求约4个月，持续在研发与客户两侧对齐功能、接口、验收标准与问题优先级，最终推动产品落地交付。",
        "因产品线覆盖较广，岗位定位偏重需求澄清、问题定位、跨部门推进与客户预期管理，而非单一机型深度专家。",
        "使用 Cursor 搭建服务部门智能知识库，沉淀常见问题、处理办法与产品文件。同事可通过对话方式检索对应文档。为销售与服务同事开通账号，部署于公司服务器并全天运行，缓解国内外时差导致的资料获取延迟。",
        "开发售后工单记录平台，支持上传视频、图片、文档，按 case 关联客户与问题，并配套数据看板统计客户问题累计，提升售后过程透明度与复盘效率。",
    ]:
        add_bullet(doc, t)

    add_job_head(doc, "ELEGOO（深圳市智能派科技有限公司）｜海外技术支持工程师", "2025.08 – 2025.10")
    add_muted(doc, "3D 打印出海｜工位远程支持欧美客户售后与技术咨询")
    for t in [
        "快速学习公司产品线，独立处理欧美客户技术咨询，形成先复现、再分层、再闭环的支持习惯。",
        "沉淀个人知识库并训练 AI Agent，将重复问题标准化，提高检索与回复效率。",
        "系统整理用户痛点、设计缺陷与改进建议，反馈产品与技术部门，作为迭代输入。",
    ]:
        add_bullet(doc, t)

    add_heading_bar(doc, "教育背景")
    add_job_head(doc, "浙江科技大学｜自动化｜本科｜GPA 3.04（专业前15%）", "2021.09 – 2025.06")
    add_muted(doc, "英语水平：CET-6｜主修单片机与嵌入式、微机原理、自动控制原理、C 语言、数据结构等｜智能车实验室三年")

    add_heading_bar(doc, "项目经历")
    add_job_head(doc, "惯导定位越野车（英飞凌 TC264 / TC377）", "2022.12 – 2023.08")
    add_muted(doc, "智能车实验室｜感知：GPS + IMU｜成果：浙江省大学生智能车竞赛二等奖（省第四）")
    for t in [
        "对比 GPS_TAU1201（UART）与 IMU660RA（SPI）方案，采用融合与互补算法提升实时性，缩短循迹响应，减少约2秒比赛用时。",
        "设计自动任务分区逻辑，降低人工干预频次。使用 VOFA+ 与 Matlab 分析路径误差并优化 PID，降低寻迹偏移。",
        "备赛过程中协调传感器选型与软硬件分工，锻炼跨角色协作与问题决策能力。",
    ]:
        add_bullet(doc, t)

    add_job_head(doc, "自动红外循迹系统（德州仪器 MSPM0G3507）", "2024.07 – 2024.08")
    add_muted(doc, "MPU6050 + 红外阵列｜成果：浙江省大学生电子设计竞赛 H 题，完成前三问")
    for t in [
        "设计 7 路红外循迹阵列（3-1-3）布局并推动团队落地。使用 DMP 获取欧拉角并建立误差修正，缓解偏航角零漂。",
        "基于 Keil5 / SysConfig 配置 GPIO 与 UART，缩短外设联调周期。",
    ]:
        add_bullet(doc, t)

    add_job_head(doc, "基于 μC/OS-II 车载电控单元（STM32F103C8T6）", "2025.02 – 2025.04")
    add_muted(doc, "实时控制与安全响应｜成果：障碍物识别响应延迟小于 50ms")
    for t in [
        "对比 FreeRTOS 与 μC/OS-II 后完成系统移植。主导超声波测距、电机减速、蜂鸣报警与霍尔转速可视化逻辑。",
        "使用嘉立创 EDA 迭代 PCB，缩短硬件验证周期，形成指标定义、方案选型、联调验证的完整闭环。",
    ]:
        add_bullet(doc, t)

    add_heading_bar(doc, "技能特长")
    for t in [
        "语言能力：英语 CET-6，可独立完成海外客户口语沟通、现场培训讲解与英文技术邮件处理。",
        "客户支持：售前培训、售后诊断、现场交付、需求转译、跨时区沟通、经销商赋能。",
        "工程基础：STM32 / 英飞凌 / TI 单片机，UART/SPI，FreeRTOS / μC/OS-II 基础，PCB 识图、焊接、万用表，Python、Matlab。",
        "效率工具：Cursor / AI Agent，知识库建设，工单与数据看板，ProcessOn / Visio / Xmind / Office。",
        "其他：持有英国签证（约2年有效）。期望工作地深圳。可接受海外出差，约一年三次，单次约一个月。",
    ]:
        add_bullet(doc, t)

    add_heading_bar(doc, "奖项荣誉")
    for t in [
        "第十八届全国大学生智能车竞赛 浙江省二等奖（2023.07）",
        "2024 年浙江省大学生电子设计竞赛 H 题，完成前三问（2024.08）",
        "浙江省大学生英语竞赛 校一等奖（2022.05）",
        "第六届普译奖全国大学生英语写作大赛 全国三等奖（2023.12）",
        "学业优秀一等奖学金、美仪企业奖学金（7/92，2022–2024）",
    ]:
        add_bullet(doc, t)

    add_heading_bar(doc, "自我评价")
    for t in [
        "沟通协调能力突出，能够在客户、销售、研发之间快速对齐事实与下一步动作，推动问题闭环。",
        "学习与工具化能力较强。面对多产品线与有限培训，能够通过自学、文档沉淀和 AI 提效形成可复制支持能力。",
        "编程能力仍在持续提升，更适合以设备理解、客户交付与跨部门推进为核心的海外技术支持或售前岗位。",
        "希望加入半导体设备或机器人公司，在深圳长期发展，并逐步向海外销售方向转型。",
    ]:
        add_bullet(doc, t)
    return doc


def build_en_1():
    doc = setup_doc((1.2, 1.2, 1.4, 1.4))
    add_center(doc, "Zhongqi Hou", size=20, bold=True, after=4)
    add_center(
        doc,
        "Target Role: Overseas Technical Support Engineer (Semiconductor Equipment / Robotics) | Shenzhen | Open to Overseas Travel",
        size=10,
        bold=True,
    )
    add_center(
        doc,
        "English-fluent with overseas customers | UK visa holder (~2 years) | +86 178-1676-9251 | zhongqi_hou@foxmail.com",
        size=9.5,
        after=4,
    )

    add_heading_bar(doc, "Summary")
    for t in [
        "Experienced in overseas field and remote support across Spain, the Netherlands, France, and Germany, with clear spoken English for customer training and troubleshooting.",
        "Strong customer–R&D coordination. Drove a key-account wheel-alignment hardware and software requirement for about four months through clarification, alignment, and product delivery.",
        "Built an AI knowledge-base assistant with Cursor for the service team, deployed on company servers around the clock. Also developed an after-sales ticket platform with media upload and issue dashboard.",
        "Automation background with smart-car lab experience in PCB reading, embedded C, soldering, and multimeter debugging. Able to ramp up on new electromechanical product lines quickly.",
    ]:
        add_bullet(doc, t, size=9.5)

    add_heading_bar(doc, "Experience")
    add_job_head(doc, "SmartSafe — Overseas Technical Support / Service Engineer", "Nov 2025 – Present")
    add_muted(doc, "Automotive aftermarket test and maintenance equipment | Pre-sales training, after-sales support, and customer–R&D coordination")
    for t in [
        "Supported product lines including wheel alignment, ADAS calibration, A/C refrigerant exchange, injector cleaning, ATF exchange, EV battery pack charge/discharge and balancing, and diagnostic tablets. Built independent support capability quickly under limited formal training.",
        "Conducted on-site fault handling and product training in Spain, the Netherlands, France, and Germany for dealers and end users.",
        "Led about four months of key-account wheel-alignment HW/SW requirement alignment between customer and R&D, clarifying scope, synchronizing validation pace, and completing delivery.",
        "Built a Q&A knowledge base with Cursor covering common fixes and product documents. Provisioned accounts for sales and service, and deployed the system on company servers around the clock to reduce timezone-related document search delays.",
        "Developed an after-sales ticket system supporting video, image, and document upload, with case-level customer linkage and a dashboard for cumulative issue tracking.",
    ]:
        add_bullet(doc, t, size=9.5)

    add_job_head(doc, "ELEGOO — Overseas Technical Support Engineer", "Aug 2025 – Oct 2025")
    add_muted(doc, "3D printing export | Remote technical support for EU and US customers")
    for t in [
        "Learned product lines quickly and independently handled EU/US technical inquiries. Built a personal knowledge base and trained an AI agent to improve response efficiency.",
        "Systematically fed recurring issues, design gaps, and improvement suggestions back to product and engineering teams.",
    ]:
        add_bullet(doc, t, size=9.5)

    add_heading_bar(doc, "Education")
    add_job_head(doc, "Zhejiang University of Science and Technology — B.E. Automation | GPA 3.04 (Top 15%)", "Sep 2021 – Jun 2025")
    for t in [
        "Three years in the Intelligent Vehicle Lab. Worked on Infineon TC264/TC377 GPS+IMU off-road vehicle with PID tuning. Provincial 2nd Prize, National Intelligent Vehicle Competition (4th in Zhejiang).",
        "Foundation in STM32, sensor drivers, UART/SPI, FreeRTOS basics, PCB reading, soldering, multimeter debugging, and embedded C.",
    ]:
        add_bullet(doc, t, size=9.5)

    add_heading_bar(doc, "Skills and Awards")
    for t in [
        "Languages: CET-6 English. Able to handle overseas customer calls, on-site training, and technical email.",
        "Support: multi-product after-sales diagnosis, pre-sales training, requirement translation, and cross-team coordination among customers, R&D, and sales.",
        "Technical: automation and embedded fundamentals, Python/Matlab data analysis, Cursor/AI tooling, knowledge-base and ticket workflow building. Prefer Shenzhen. Open to about three overseas trips per year, about one month each.",
        "Awards: Provincial 2nd Prize, National Intelligent Vehicle Competition (2023). Campus 1st Prize, Zhejiang College English Contest (2022). National 3rd Prize, Puyi English Writing Contest (2023). Scholarships (rank 7/92).",
    ]:
        add_bullet(doc, t, size=9.5)

    add_heading_bar(doc, "Profile")
    add_line(
        doc,
        "Strong communicator and fast learner, effective at clarifying field issues and driving solutions between customers and R&D. Seeking to deepen overseas technical support in semiconductor equipment or robotics, then grow toward overseas sales.",
        size=9.5,
    )
    return doc


def build_en_2():
    doc = setup_doc((1.4, 1.4, 1.5, 1.5))
    add_center(doc, "Zhongqi Hou", size=20, bold=True, after=4)
    add_center(
        doc,
        "Target Role: Overseas Technical Support Engineer (Semiconductor Equipment / Robotics) | Shenzhen",
        size=10.5,
        bold=True,
    )
    add_center(doc, "English-fluent with overseas customers | UK visa holder (~2 years)", size=10)
    add_center(
        doc,
        "+86 178-1676-9251 | zhongqi_hou@foxmail.com | Travel: about 3 overseas trips/year, about 1 month each",
        size=10,
        after=4,
    )

    add_heading_bar(doc, "Professional Summary")
    for t in [
        "Career path: deepen overseas technical support first, then move toward overseas sales or pre-sales roles.",
        "Current strength: independent overseas customer delivery through remote support and Europe on-site work. Able to translate field problems into R&D-executable requirements and communicate solutions back to customers clearly.",
        "Efficiency edge: use Cursor to build department-level knowledge systems and ticket tooling that reduce timezone friction and improve case reuse.",
        "Technical base: Automation degree and Intelligent Vehicle Lab experience covering sensors, embedded systems, PCB reading, and hands-on debugging for electromechanical and robotics-related equipment support.",
    ]:
        add_bullet(doc, t)

    add_heading_bar(doc, "Work Experience")
    add_job_head(doc, "SmartSafe — Overseas Technical Support / Service Engineer", "Nov 2025 – Present")
    add_muted(doc, "Automotive aftermarket inspection and maintenance equipment | Internal service engineer and external overseas technical support")
    for t in [
        "Supported multi-line products including wheel alignment, ADAS calibration, A/C refrigerant exchange, injector cleaning, transmission fluid exchange, EV battery pack charge/discharge and balancing, and diagnostic tablets. Built deliverable support capability quickly despite limited deep product training.",
        "Completed overseas field work in Spain, the Netherlands, France, and Germany for after-sales troubleshooting, product training, and customer Q&A with dealers and end users.",
        "Followed key-account wheel-alignment hardware and software requirements for about four months, continuously aligning features, interfaces, acceptance criteria, and issue priorities between R&D and the customer until product delivery.",
        "Role focus: reliable at requirement clarification, issue triage, cross-team execution, and customer expectation management across a wide portfolio.",
        "Built an AI knowledge-base assistant with Cursor for the service team, covering common issues, fix methods, and product documents with chat-style retrieval. Provisioned accounts for sales and service, hosted on company servers around the clock to reduce after-hours document search across time zones.",
        "Developed an after-sales ticket platform with video, image, and document upload, customer–issue linkage, and a dashboard for cumulative issue visibility and review.",
    ]:
        add_bullet(doc, t)

    add_job_head(doc, "ELEGOO — Overseas Technical Support Engineer", "Aug 2025 – Oct 2025")
    add_muted(doc, "3D printing export | Desk-based remote after-sales and technical support for EU and US customers")
    for t in [
        "Ramped on product lines quickly and independently handled EU/US technical inquiries. Built a personal knowledge base and trained an AI agent to standardize recurring issues and accelerate response.",
        "Compiled user pain points, design defects, and improvement suggestions as structured input for product and engineering teams.",
    ]:
        add_bullet(doc, t)

    add_heading_bar(doc, "Education")
    add_job_head(doc, "Zhejiang University of Science and Technology — B.E. Automation | GPA 3.04 (Top 15%)", "Sep 2021 – Jun 2025")
    add_muted(doc, "CET-6 | Coursework: MCU/embedded design, microcomputer principles, automatic control, C programming, data structures | Intelligent Vehicle Lab for 3 years")

    add_heading_bar(doc, "Projects")
    add_job_head(doc, "Inertial Navigation Off-Road Vehicle (Infineon TC264 / TC377)", "Dec 2022 – Aug 2023")
    add_muted(doc, "GPS + IMU | Provincial 2nd Prize (4th in Zhejiang), National Intelligent Vehicle Competition")
    for t in [
        "Compared GPS over UART and IMU over SPI, then used fusion and complementary filtering to improve real-time response and reduce race time by about two seconds.",
        "Designed automatic task-zone logic to reduce manual intervention. Used VOFA+ and Matlab to analyze path error and tune PID. Coordinated sensor selection and HW/SW planning during competition prep.",
    ]:
        add_bullet(doc, t)

    add_job_head(doc, "Infrared Line-Following (TI MSPM0G3507) and Vehicle ECU on μC/OS-II (STM32F103)", "2024.07 – 2025.04")
    for t in [
        "Designed IR array plus MPU6050 sensing with DMP Euler angles and yaw-drift correction. Zhejiang Electronic Design Contest Problem H, first three tasks completed.",
        "Ported μC/OS-II after comparing with FreeRTOS. Owned ultrasonic ranging, motor deceleration, buzzer alert, and Hall RPM display. Achieved obstacle response under 50 ms and iterated PCB in JLCPCB EDA.",
    ]:
        add_bullet(doc, t)

    add_heading_bar(doc, "Skills")
    for t in [
        "Languages: CET-6 English for overseas customer conversations, on-site training, and technical email.",
        "Support: pre-sales training, after-sales diagnosis, field delivery, requirement translation, and cross-timezone communication.",
        "Engineering: STM32 / Infineon / TI MCUs, UART/SPI, RTOS basics, PCB reading, soldering, multimeter, Python, Matlab.",
        "Productivity: Cursor / AI agents, knowledge-base building, ticket and dashboard workflows, ProcessOn / Visio / Xmind / Office.",
        "Other: UK visa (~2 years). Prefer Shenzhen base. Overseas travel acceptable, about three times per year, about one month each.",
    ]:
        add_bullet(doc, t)

    add_heading_bar(doc, "Honors")
    for t in [
        "Provincial 2nd Prize, 18th National College Student Intelligent Vehicle Competition (Jul 2023)",
        "Campus 1st Prize, Zhejiang College English Contest (May 2022)",
        "National 3rd Prize, 6th Puyi National College English Writing Contest (Dec 2023)",
        "Academic Excellence and Meiyi Corporate Scholarship (rank 7/92, 2022–2024)",
    ]:
        add_bullet(doc, t)

    add_heading_bar(doc, "Self Assessment")
    add_line(
        doc,
        "Communication and cross-team coordination are current strengths. Fast learner with a tooling mindset. Coding depth is still improving. Best fit is overseas technical support or pre-sales in semiconductor equipment or robotics, based in Shenzhen, with a medium-term path toward overseas sales.",
        size=10,
    )
    return doc


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    files = {
        "侯钟棋-海外技术支持-中文-一页版.docx": build_zh_1(),
        "侯钟棋-海外技术支持-中文-两页版.docx": build_zh_2(),
        "Hou-Zhongqi-Overseas-Technical-Support-EN-1page.docx": build_en_1(),
        "Hou-Zhongqi-Overseas-Technical-Support-EN-2page.docx": build_en_2(),
    }
    for name, doc in files.items():
        path = OUT / name
        doc.save(path)
        print("Wrote", path)


if __name__ == "__main__":
    main()
